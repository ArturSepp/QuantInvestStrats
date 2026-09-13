"""Render and serialize completed stress results, with an optional parser appendix."""

from dataclasses import dataclass, field
import hashlib
import json
import re
from pathlib import Path
from types import MappingProxyType
from typing import Mapping

import numpy as np
import pandas as pd

from qis.portfolio.stress.analytics import PortfolioStressResult
from qis.portfolio.stress._clusters import compute_cluster_contributions


@dataclass(frozen=True)
class StressReportConfig:
    """Presentation options and caller-supplied estimation diagnostics.

    Attributes:
        title: Report heading and PDF title.
        model_label: Optional plain model/version label.
        factor_labels: Display aliases, keyed by actual factor ID.
        response_diagnostics: Optional response-indexed diagnostics, including r2.
        cluster_memberships: Optional fitted group-to-response membership Series.
        cluster_linkages: Matching fitted linkage arrays in membership index order.
        cluster_cutoffs: Matching fitted cutoffs; never estimated by the report.
        selected_grids: Up to six caller-named grids for the sensitivity page.
            Empty selects the ranked individual factors or first six custom grids;
            all grids are always exported.
        notes: Plain methodology/coverage notes supplied by the application.
        write_workbook: Write a numerical workbook using the existing QIS serializer.
        write_previews: Also save PNG previews for inspection.
        model_name: Short model name used in slide and figure titles.
        appendix_table: Optional preformatted parser-owned table for page eleven.
        appendix_title: Parser-owned page title.
        appendix_subtitle: Parser-owned description above the table.
        appendix_notes: Parser-owned variable definitions and source explanations.
        report_name: Optional account/portfolio name on the first page; defaults to title.
            The header appends the actual notional, currency and assigned model label.
    """

    title: str = "Portfolio stress report"
    model_label: str = ""
    factor_labels: Mapping[str, str] = field(default_factory=dict)
    response_diagnostics: pd.DataFrame | None = None
    cluster_memberships: Mapping[str, pd.Series] = field(default_factory=dict)
    cluster_linkages: Mapping[str, np.ndarray] = field(default_factory=dict)
    cluster_cutoffs: Mapping[str, float] = field(default_factory=dict)
    selected_grids: tuple[str, ...] = ()
    notes: tuple[str, ...] = ()
    write_workbook: bool = True
    write_previews: bool = False
    model_name: str = "Factor model"
    appendix_table: pd.DataFrame | None = None
    appendix_title: str = "Coverage and estimation quality"
    appendix_subtitle: str = ""
    appendix_notes: tuple[str, ...] = ()
    report_name: str | None = None

    def __post_init__(self):
        """Snapshot caller diagnostics and validate presentation choices."""
        if self.report_name is not None:
            if not isinstance(self.report_name, str) or not self.report_name.strip():
                raise ValueError("report_name must be a nonempty string or None")
            object.__setattr__(self, "report_name", self.report_name.strip())
        if not self.title:
            raise ValueError("report title must be nonempty")
        if not self.model_name or not self.appendix_title:
            raise ValueError("model_name and appendix_title must be nonempty")
        if self.appendix_table is not None:
            if self.appendix_table.empty:
                raise ValueError("appendix_table must be nonempty or None")
            if len(self.appendix_table) > 24 or len(self.appendix_table.columns) > 10:
                raise ValueError("appendix_table supports at most 24 rows and 10 columns")
            object.__setattr__(self, "appendix_table", self.appendix_table.copy(deep=True))
        object.__setattr__(self, "appendix_notes", tuple(self.appendix_notes))
        grids = tuple(self.selected_grids)
        if len(grids) > 6 or len(set(grids)) != len(grids):
            raise ValueError("selected_grids must contain at most six unique names")
        if not (
            set(self.cluster_memberships) == set(self.cluster_linkages) == set(self.cluster_cutoffs)
        ):
            raise ValueError("cluster membership, linkage and cutoff keys must agree")
        if self.response_diagnostics is not None:
            diagnostics = self.response_diagnostics.copy(deep=True)
            if diagnostics.index.has_duplicates or diagnostics.columns.has_duplicates:
                raise ValueError("response diagnostics require unique labels")
            object.__setattr__(self, "response_diagnostics", diagnostics)
        object.__setattr__(self, "factor_labels", MappingProxyType(dict(self.factor_labels)))
        object.__setattr__(self, "selected_grids", grids)
        object.__setattr__(self, "notes", tuple(self.notes))
        object.__setattr__(
            self,
            "cluster_memberships",
            MappingProxyType(
                {key: value.copy(deep=True) for key, value in self.cluster_memberships.items()}
            ),
        )
        object.__setattr__(
            self,
            "cluster_linkages",
            MappingProxyType(
                {
                    key: np.asarray(value, dtype=float).copy()
                    for key, value in self.cluster_linkages.items()
                }
            ),
        )
        object.__setattr__(self, "cluster_cutoffs", MappingProxyType(dict(self.cluster_cutoffs)))


@dataclass(frozen=True)
class StressReportArtifacts:
    """Paths written by a report operation; the numerical result remains reusable.

    Attributes:
        pdf_path: Ten core pages, plus page eleven when a parser supplies a table.
        table_paths: Numerical table names mapped to CSV paths.
        workbook_path: Optional workbook path.
        manifest_path: JSON with conventions, table mapping and content hashes.
        preview_paths: Optional page PNGs.
    """

    pdf_path: Path
    table_paths: Mapping[str, Path]
    workbook_path: Path | None
    manifest_path: Path
    preview_paths: tuple[Path, ...] = ()



def _report_heading(result, config):
    """Describe the account, actual reporting notional and assigned risk model."""
    meta = result.metadata
    name = config.report_name or config.title
    model = config.model_label or config.model_name
    return (f"{name} | Notional {meta['reference_currency']} "
            f"{meta['reporting_denominator']:,.0f} | Risk model {model}")


def _loading_table(result, config):
    """Join copied model diagnostics and absolute-exposure-weighted fit summaries."""
    order = result.response_exposures.abs().sort_values(ascending=False, kind="stable").index
    table = result.factor_loadings.join(result.report_diagnostics["Unit response risk"])
    r2 = (
        config.response_diagnostics["r2"].reindex(order)
        if config.response_diagnostics is not None and "r2" in config.response_diagnostics
        else pd.Series(np.nan, index=order)
    )
    table["R-squared"] = r2
    table = table.loc[order[:20]].copy()
    aggregates = result.report_diagnostics["Loading aggregates"].copy()
    for label in aggregates.index:
        ids = order[20:] if label == "Rest of assets" else order
        available = r2.loc[ids].dropna()
        weights = result.response_exposures.reindex(available.index).abs()
        aggregates.loc[label, "R-squared"] = (
            available @ weights / weights.sum() if weights.sum() else np.nan
        )
    table["response_exposure"] = result.response_exposures
    table = pd.concat([table, aggregates])
    return table.loc[
        :,
        [
            *result.factor_loadings.columns,
            "R-squared",
            "Model total vol",
            "Systematic vol",
            "Idio vol",
            "response_exposure",
        ],
    ]


def _report_tables(result, config):
    """Collect all numerical observations without the PDF's display row limits."""
    tables = dict(result.report_diagnostics)
    tables["Displayed loadings and fit"] = _loading_table(result, config)
    for key, summary in result.summaries.items():
        tables[f"{key} summary"] = summary
    tables["Historical worst months"] = result.historical_ranking
    tables["Current factor exposures"] = pd.concat(
        [result.factor_exposures, result.factor_betas], axis=1
    )
    tables["Current risk"] = result.risk.to_frame("value")
    tables["Holding factor exposures"] = result.holding_factor_exposures
    tables["Factor group exposures"] = result.factor_group_exposures
    tables["Holding local risk"] = result.holding_risk
    tables["Response risk contributions"] = result.response_risk_contributions
    for key, value in result.valuations.items():
        tables[f"{key} factor log shocks"] = value.factor_log_shocks
        tables[f"{key} holding pnl"] = value.pnl
        tables[f"{key} holding mtm"] = value.mtm
        tables[f"{key} attribution"] = result.attribution[key]
    for key, value in result.grids.items():
        tables[f"Grid {key} summary"] = result.grid_summaries[key]
        tables[f"Grid {key} log shocks"] = value.factor_log_shocks
        tables[f"Grid {key} holding pnl"] = value.pnl
        tables[f"Grid {key} holding mtm"] = value.mtm
    if result.historical is not None:
        tables["Historical all factor shocks"] = result.historical.factor_log_shocks
        tables["Historical all holding pnl"] = result.historical.pnl
        tables["Historical all holding mtm"] = result.historical.mtm
    tables["Positions and payoff audit"] = result.positions
    tables["Vanilla leg terms"] = result.leg_terms
    tables["Holding response Jacobian"] = result.response_jacobian
    tables["Shared response exposures"] = result.response_exposures.to_frame()
    tables["Underlying response betas"] = result.factor_loadings
    tables["Annual factor covariance"] = result.factor_covariance
    tables["Annual residual variances"] = result.residual_variances.to_frame("variance")
    tables["Historical coverage"] = result.historical_coverage
    tables["Grid conventions"] = result.grid_metadata
    if config.response_diagnostics is not None:
        tables["Supplied fit diagnostics"] = config.response_diagnostics
    for key, members in config.cluster_memberships.items():
        tables[f"Cluster {key} membership"] = members.to_frame("cluster")
        tables[f"Cluster {key} linkage"] = pd.DataFrame(
            config.cluster_linkages[key], columns=["left", "right", "distance", "count"]
        )
    clusters = compute_cluster_contributions(result, config.cluster_memberships)
    tables["Cluster holding assignments"] = clusters.holdings
    tables["Cluster portfolio summary"] = clusters.summary
    tables["Cluster weighted factor exposures"] = clusters.factor_exposures
    tables["Cluster dollar factor exposures"] = clusters.factor_dollars
    tables["Cluster Euler risk contributions"] = clusters.risk
    tables["Cluster display grouping"] = clusters.display_groups.to_frame()
    for name in clusters.scenario_pnl:
        tables[f"Cluster {name} pnl"] = clusters.scenario_pnl[name]
        tables[f"Cluster {name} NAV contributions"] = clusters.scenario_nav[name]
    if config.appendix_table is not None:
        tables["Parser appendix"] = config.appendix_table
    metadata = dict(result.metadata)
    metadata.update(
        {
            "title": config.title,
            "report_name": config.report_name or config.title,
            "report_heading": _report_heading(result, config),
            "model_label": config.model_label,
            "report_notes": list(config.notes),
            "appendix_notes": list(config.appendix_notes),
        }
    )
    tables["Conventions"] = pd.DataFrame(
        {"value": {key: json.dumps(value, ensure_ascii=False) for key, value in metadata.items()}}
    )
    return tables


def _format_workbook(path):
    """Format exported tables without changing their stored numerical values."""
    from datetime import date, datetime
    from math import ceil
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="183A50")
    stripe_fill = PatternFill("solid", fgColor="EEF3F7")
    percent_columns = {
        "annual_vol", "euler_vol", "variance_share", "portfolio_return", "r2", "rsquared",
        "r_squared", "R-squared", "annual_total_vol", "annual_systematic_vol",
        "annual_residual_vol", "annual_factor_model_vol", "lower_bound", "upper_bound",
        "band_half_width",
    }
    for sheet in book:
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 44
        if not sheet.merged_cells.ranges:
            sheet.auto_filter.ref = sheet.dimensions
        headers = {cell.column: str(cell.value or "") for cell in sheet[1]}
        sheet_percent_columns = percent_columns
        has_grid_audit = headers.get(1) == "grid" and headers.get(2) == "factor_return"
        pane = "C2" if has_grid_audit else "B2"
        if sheet.freeze_panes != pane:
            sheet.freeze_panes = pane
        # Reapplying openpyxl freeze_panes appends duplicate selections; keep one per pane.
        selections = {}
        for selection in sheet.sheet_view.selection:
            selection.activeCell = selection.activeCell or "A1"
            selection.sqref = selection.sqref or "A1"
            selections.setdefault(selection.pane, selection)
        sheet.sheet_view.selection = list(selections.values())
        sheet_percent_columns = sheet_percent_columns | {
            "lower_1sigma", "upper_1sigma", "lower_2sigma", "upper_2sigma",
            "conditional_vol_horizon", "conditional_factor_vol_horizon", "residual_vol_horizon",
        }
        if has_grid_audit:
            sheet_percent_columns = percent_columns | {
                "factor_return", "mean", "mean_se", "mean_ci_lower", "mean_ci_upper", "confidence",
            }
            for merged in list(sheet.merged_cells.ranges):
                if merged.min_col == merged.max_col == 1:
                    grid = sheet.cell(merged.min_row, 1).value
                    first, last = merged.min_row, merged.max_row
                    sheet.unmerge_cells(str(merged))
                    for row_number in range(first, last + 1):
                        sheet.cell(row_number, 1, grid)
            sheet.auto_filter.ref = sheet.dimensions
            if "Grid conditional" in sheet.title:
                sheet_percent_columns |= set(headers.values()) - {"grid"}
        for column in range(1, sheet.max_column + 1):
            sample = [sheet.cell(row, column).value for row in range(2, min(sheet.max_row, 80) + 1)]
            text_width = max((len(str(value)) for value in sample
                              if isinstance(value, str)), default=0)
            width = max(18, min(48, text_width + 2), min(28, len(headers[column]) + 2))
            sheet.column_dimensions[get_column_letter(column)].width = (
                max(32, width) if column == 1 else width
            )
        for row in sheet:
            height = 18
            for cell in row:
                if isinstance(cell, MergedCell):
                    continue
                cell.font = Font(name="Calibri", size=11, color="183A50")
                cell.alignment = Alignment(vertical="center")
                if cell.row == 1:
                    cell.fill = header_fill
                    cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                    continue
                if cell.row % 2 == 0:
                    cell.fill = stripe_fill
                if isinstance(cell.value, (datetime, date)):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, float):
                    cell.number_format = (
                        "0.00%;[Red](0.00%);0.00%"
                        if headers[cell.column] in sheet_percent_columns
                        else "#,##0.0000;[Red](#,##0.0000);0.0000"
                    )
                elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                    cell.number_format = (
                        "0.00%;[Red](0.00%);0.00%" if headers[cell.column] in sheet_percent_columns
                        else "0" if cell.column == 1 or headers[cell.column].endswith("_id")
                        else "#,##0;[Red](#,##0);0"
                    )
                elif isinstance(cell.value, str):
                    width = sheet.column_dimensions[cell.column_letter].width
                    if len(cell.value) <= 400:
                        cell.alignment = Alignment(wrap_text=True, vertical="center")
                        height = max(height, 16 * ceil(len(cell.value) / max(width - 3, 1)))
            if row[0].row > 1:
                sheet.row_dimensions[row[0].row].height = min(height, 180)
        if sheet.title == "Contents":
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, 1)
                if cell.value in book.sheetnames:
                    cell.hyperlink = "#'" + str(cell.value).replace("'", "''") + "'!A1"
                    cell.font = Font(name="Calibri", size=11, color="1264A3", underline="single")
    book.save(path)


def generate_portfolio_stress_report(
    result: PortfolioStressResult, output_dir: str | Path, config: StressReportConfig | None = None
) -> StressReportArtifacts:
    """Render a completed result without fitting or re-evaluating any payoff.

    The PDF has ten core subjects and an optional parser-supplied eleventh page. Display limits are
    explicitly labelled; CSV/workbook tables retain every holding/scenario/grid.

    Args:
        result: Detached result returned by run_portfolio_stress_test.
        output_dir: Fresh output directory; existing directories are never overwritten.
        config: Titles, display diagnostics and artifact preferences.

    Returns:
        Paths to the PDF, CSV tables, optional workbook and audit manifest.
    """
    config = config or StressReportConfig()
    if set(config.selected_grids) - set(result.grids):
        raise ValueError("selected_grids includes a grid absent from the result")
    if set(config.factor_labels) - set(result.factor_loadings.columns):
        raise ValueError("factor_labels includes an unknown fitted factor")
    if config.response_diagnostics is not None:
        if not set(config.response_diagnostics.index).issubset(result.factor_loadings.index):
            raise ValueError("response diagnostics includes an unknown fitted response")
    for members in config.cluster_memberships.values():
        if not set(members.index).issubset(result.factor_loadings.index):
            raise ValueError("cluster membership includes an unknown fitted response")
    from matplotlib.backends.backend_pdf import PdfPages
    import matplotlib.pyplot as plt
    from qis.portfolio.stress._figures import report_pages

    output_dir = Path(output_dir).expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=False)
    pdf_path = output_dir / "portfolio_stress_report.pdf"
    tables = _report_tables(result, config)
    table_dir = output_dir / "tables"
    table_dir.mkdir()
    table_paths = {}
    table_manifest = []
    workbook_tables = {}
    for number, (name, frame) in enumerate(tables.items(), 1):
        path = table_dir / f"{number:02d}.csv"
        frame.to_csv(path)
        table_paths[name] = path
        prefix = f"{number:02d} "
        clean_name = re.sub(r"[\\/?*:\[\]]", "_", name)
        sheet = prefix + clean_name[: 31 - len(prefix)]
        workbook_tables[sheet] = frame
        table_manifest.append(
            {
                "name": name,
                "workbook_sheet": sheet,
                "csv": str(path.relative_to(output_dir)),
                "rows": len(frame),
                "columns": list(map(str, frame.columns)),
            }
        )
    workbook_path = None
    if config.write_workbook:
        # Use the owning QIS serialization API; report rendering owns no Excel backend.
        from qis.file_utils import save_df_dict_to_excel

        workbook_path = Path(
            save_df_dict_to_excel(
                {
                    "Contents": pd.DataFrame(table_manifest).set_index("workbook_sheet"),
                    **workbook_tables,
                },
                file_name="portfolio_stress_tables",
                local_path=str(output_dir),
            )
        )
        _format_workbook(workbook_path)
    previews, titles = [], []
    with PdfPages(
        pdf_path,
        metadata={
            "Title": config.title,
            "Author": "QIS",
            "Subject": "Intrinsic portfolio factor stress",
        },
    ) as pdf:
        for number, (title, figure) in enumerate(report_pages(result, config), 1):
            try:
                pdf.savefig(figure)
                titles.append(title)
                if config.write_previews:
                    path = output_dir / f"page_{number:02d}.png"
                    figure.savefig(path, dpi=120, facecolor="white")
                    previews.append(path)
            finally:
                plt.close(figure)
    all_paths = [pdf_path, *table_paths.values(), *previews]
    if workbook_path is not None:
        all_paths.append(workbook_path)
    manifest = {
        "metadata": dict(result.metadata),
        "title": config.title,
        "report_name": config.report_name or config.title,
        "report_heading": _report_heading(result, config),
        "model_label": config.model_label,
        "report_notes": list(config.notes),
        "page_count": len(titles),
        "page_titles": titles,
        "tables": table_manifest,
        "display_limits": {
            "scenario_rows": 12,
            "contributors": 10,
            "response_rows": 20,
            "response_factor_columns": len(result.factor_loadings.columns),
            "factor_panels": 6,
            "grid_panels": 4,
        },
        "hashes": {
            str(path.relative_to(output_dir)): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in all_paths
        },
    }
    manifest_path = output_dir / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    return StressReportArtifacts(
        pdf_path, MappingProxyType(table_paths), workbook_path, manifest_path, tuple(previews)
    )
